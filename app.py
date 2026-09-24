import json, os, re, secrets, sqlite3, time
from datetime import timedelta
from io import BytesIO
from xml.sax.saxutils import escape
from flask import Flask, abort, g, jsonify, redirect, request, send_file, send_from_directory, session
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__, static_folder="static")
DB = os.environ.get("DB_PATH", os.path.join(app.root_path, "cricket.db"))
def init():
    c = sqlite3.connect(DB)
    c.execute("CREATE TABLE IF NOT EXISTS kv (k TEXT PRIMARY KEY, v TEXT NOT NULL)")
    c.execute("CREATE TABLE IF NOT EXISTS matches (id INTEGER PRIMARY KEY, data TEXT NOT NULL, updated REAL NOT NULL)")
    c.execute("CREATE TABLE IF NOT EXISTS users (username TEXT PRIMARY KEY, pw TEXT NOT NULL, role TEXT NOT NULL)")
    r = c.execute("SELECT v FROM kv WHERE k='secret'").fetchone()
    key = r[0] if r else secrets.token_hex(32)
    if not r:
        c.execute("INSERT INTO kv VALUES('secret',?)", (key,))
    u, p = os.environ.get("SCORER_USER", "scorer"), os.environ.get("SCORER_PASS")
    if p or not c.execute("SELECT 1 FROM users WHERE role='scorer'").fetchone():
        gen = not p
        p = p or secrets.token_urlsafe(8)
        c.execute("INSERT OR REPLACE INTO users VALUES(?,?,?)", (u, generate_password_hash(p), "scorer"))
        if gen:
            print("[setup] scorer login ->", u, "/", p, flush=True)
    c.commit()
    c.close()
    return key


app.secret_key = init()
app.config.update(SESSION_COOKIE_SAMESITE="Lax", PERMANENT_SESSION_LIFETIME=timedelta(days=30),
                  SEND_FILE_MAX_AGE_DEFAULT=0, SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE") == "1")


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB)
        g.db.execute("CREATE TABLE IF NOT EXISTS matches "
                     "(id INTEGER PRIMARY KEY, data TEXT NOT NULL, updated REAL NOT NULL)")
        g.db.execute("CREATE TABLE IF NOT EXISTS kv (k TEXT PRIMARY KEY, v TEXT NOT NULL)")
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()


def guard():
    if session.get("role") != "scorer":
        abort(403)


@app.get("/")
def index():
    with open(os.path.join(app.root_path, "static", "index.html"), encoding="utf-8") as f:
        html = f.read()
    return html.replace("__ROLE__", session["role"]).replace("__USER__", session["u"])


@app.get("/api/matches")
def list_matches():
    rows = db().execute("SELECT data FROM matches ORDER BY id DESC").fetchall()
    return jsonify([json.loads(r[0]) for r in rows])


@app.get("/api/matches/<int:mid>")
def get_match(mid):
    row = db().execute("SELECT data FROM matches WHERE id=?", (mid,)).fetchone()
    if not row:
        abort(404)
    return jsonify(json.loads(row[0]))


@app.put("/api/matches/<int:mid>")
def put_match(mid):
    guard()
    m = request.get_json(silent=True)
    if not isinstance(m, dict) or m.get("id") != mid:
        abort(400)
    d = db()
    d.execute("INSERT INTO matches(id,data,updated) VALUES(?,?,?) "
              "ON CONFLICT(id) DO UPDATE SET data=excluded.data, updated=excluded.updated",
              (mid, json.dumps(m), time.time()))
    d.commit()
    return jsonify(ok=True)


@app.delete("/api/matches/<int:mid>")
def delete_match(mid):
    guard()
    d = db()
    d.execute("DELETE FROM matches WHERE id=?", (mid,))
    d.commit()
    return jsonify(ok=True)


@app.get("/api/teams")
def get_teams():
    r = db().execute("SELECT v FROM kv WHERE k='teams'").fetchone()
    return jsonify(json.loads(r[0]) if r else [])


@app.put("/api/teams")
def put_teams():
    guard()
    t = request.get_json(silent=True)
    if not isinstance(t, list):
        abort(400)
    d = db()
    d.execute("INSERT INTO kv(k,v) VALUES('teams',?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (json.dumps(t),))
    d.commit()
    return jsonify(ok=True)


def _pad(rows):
    n = max((len(r) for r in rows), default=0)
    return [["" if c is None else c for c in r] + [""] * (n - len(r)) for r in rows]


def _head(r, i):
    return i == 0 or (r and str(r[0]).startswith("\u00bb"))


@app.post("/api/export/<fmt>")
def export(fmt):
    d = request.get_json(silent=True) or {}
    title, sheets = str(d.get("title", "Tournament")), d.get("sheets", [])
    buf = BytesIO()
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        wb.remove(wb.active)
        for sh in sheets:
            ws = wb.create_sheet(str(sh["name"])[:31])
            rows = _pad(sh["rows"])
            for r in rows:
                ws.append(r)
            for i, r in enumerate(rows):
                if _head(r, i):
                    for c in ws[i + 1]:
                        c.font = Font(bold=True)
            if rows:
                for c in ws[1]:
                    c.fill = PatternFill("solid", fgColor="DDEEDD")
            for col in ws.columns:
                w = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), 60)
        wb.save(buf)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        ss = getSampleStyleSheet()
        cell = ParagraphStyle("c", parent=ss["BodyText"], fontSize=8, leading=10)
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                                topMargin=28, bottomMargin=28, title=title)
        W = landscape(A4)[0] - 56
        el = [Paragraph(escape(title), ss["Title"])]
        for sh in sheets:
            rows = _pad(sh["rows"])
            if not rows or not rows[0]:
                continue
            el += [Spacer(1, 8), Paragraph(escape(str(sh["name"])), ss["Heading2"])]
            lens = [min(max(len(str(r[i])) for r in rows), 40) + 4 for i in range(len(rows[0]))]
            data = []
            for i, r in enumerate(rows):
                data.append([Paragraph(("<b>%s</b>" if _head(r, i) else "%s") % escape(str(c)), cell) for c in r])
            t = Table(data, colWidths=[W * l / sum(lens) for l in lens], repeatRows=1)
            st = [("GRID", (0, 0), (-1, -1), .25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEEDD"))]
            st += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F0F0F0"))
                   for i, r in enumerate(rows) if i and _head(r, i)]
            t.setStyle(TableStyle(st))
            el.append(t)
        doc.build(el)
        mime = "application/pdf"
    else:
        abort(404)
    buf.seek(0)
    return send_file(buf, mimetype=mime, as_attachment=True, download_name="tournament." + fmt)


@app.before_request
def gate():
    p = request.path
    if p.startswith("/static/") or p in ("/login", "/api/login", "/logout") or "u" in session:
        return
    if p.startswith("/api/"):
        abort(401)
    return redirect("/login")


@app.get("/login")
def login_page():
    return send_from_directory("static", "login.html")


@app.get("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.post("/api/login")
def login():
    d = request.get_json(silent=True) or {}
    u = str(d.get("username", "")).strip()
    r = db().execute("SELECT pw, role FROM users WHERE username=?", (u,)).fetchone()
    if not r or not check_password_hash(r[0], str(d.get("password", ""))):
        time.sleep(0.8)
        return jsonify(error="Wrong username or password"), 401
    if d.get("role") and d["role"] != r[1]:
        return jsonify(error="This is a %s account - use the %s tab" % (r[1], r[1])), 403
    session.clear()
    session.permanent = True
    session["u"], session["role"] = u, r[1]
    return jsonify(ok=True)


@app.get("/api/users")
def list_users():
    guard()
    rows = db().execute("SELECT username, role FROM users ORDER BY role, username").fetchall()
    return jsonify([{"username": u, "role": r} for u, r in rows])


@app.put("/api/users")
def save_user():
    guard()
    d = request.get_json(silent=True) or {}
    u, p, role = str(d.get("username", "")).strip(), str(d.get("password", "")), d.get("role")
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,30}", u):
        return jsonify(error="Username: 3-30 letters, numbers, . _ -"), 400
    if len(p) < 6:
        return jsonify(error="Password must be at least 6 characters"), 400
    if role not in ("scorer", "viewer") or (u == session["u"] and role != "scorer"):
        return jsonify(error="Invalid role"), 400
    c = db()
    c.execute("INSERT OR REPLACE INTO users VALUES(?,?,?)", (u, generate_password_hash(p), role))
    c.commit()
    return jsonify(ok=True)


@app.delete("/api/users/<name>")
def del_user(name):
    guard()
    if name == session["u"]:
        return jsonify(error="You cannot delete your own account"), 400
    c = db()
    c.execute("DELETE FROM users WHERE username=?", (name,))
    c.commit()
    return jsonify(ok=True)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
